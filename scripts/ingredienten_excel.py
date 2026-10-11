#!/usr/bin/env python3
"""
Maakt van ingredienten_overzicht.json (uit `npm run ingredienten`) een Excel
om ingrediëntnamen samen te voegen, zodat meer regels van de lijst in het
AH-mandje landen in plaats van als zoeklink.

    npm run ingredienten -- --offline
    python3 scripts/ingredienten_excel.py            # schrijft ingredienten.xlsx in de root

Tabbladen:
  Leeswijzer          wat er in staat en hoe je het invult
  Zoeklinks           elk ingrediënt dat nu een zoeklink krijgt, met een voorstel
                      waar het naartoe kan en lege kolommen voor je besluit
  Dubbele lijstregels namen die hetzelfde AH-product krijgen maar op de lijst
                      als aparte regels staan
  Alle ingrediënten   elke naam uit de recepten met wat de app ermee doet
  AH-mapping          elke mapping-regel met hoe vaak hij geraakt wordt

De voorstellen zijn heuristiek (zelfde lijstregel, keuze "x of y", deel van
de naam, gelijkend woord). Ze zijn bedoeld om na te kijken, niet om blind
over te nemen: een verkeerd product in het mandje merk je pas bij de kassa.

Vereist openpyxl (pip install openpyxl).
"""

from __future__ import annotations

import difflib
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:  # pragma: no cover
    sys.exit('openpyxl ontbreekt: pip install openpyxl')

ROOT = Path(__file__).resolve().parent.parent
BRON = ROOT / 'ingredienten_overzicht.json'
DOEL = ROOT / 'ingredienten.xlsx'

FONT = 'Arial'
KOP = Font(name=FONT, bold=True, color='FFFFFF')
KOPVULLING = PatternFill('solid', fgColor='8B1E1E')  # Pinch-rood
ROOD = PatternFill('solid', fgColor='FADBD8')
ORANJE = PatternFill('solid', fgColor='FDEBD0')
GEEL = PatternFill('solid', fgColor='FFF9C4')
GROEN = PatternFill('solid', fgColor='E8F5E9')
INVUL = PatternFill('solid', fgColor='FFFDE7')


def plat(tekst: str) -> str:
    """Zoals ingredientKey, grofweg: kleine letters, zonder accenten en leestekens."""
    t = unicodedata.normalize('NFD', tekst.lower())
    t = ''.join(c for c in t if not unicodedata.combining(c))
    t = re.sub(r"'s\b", ' ', t)
    t = re.sub(r'\([^)]*\)', ' ', t)
    t = re.sub(r'[^a-z ]', ' ', t)
    return re.sub(r'\s+', ' ', t).strip()


# Woorden die het product niet veranderen als ze wegvallen; aangevuld met wat de
# app in zoekProduct als ONSCHULDIG kent. Alleen voor het voorstel, niet voor de app.
ONSCHULDIG = set("""
gewone gewoon normale normaal
in of en van voor op om te een plus de met erbij naar ter bij het je aan dan als ook
blokje blokjes plakje plakjes reepje reepjes ring ringen stuk stukken stukje stukjes stuks
part partje partjes parten linten vieren lengte dun dunne dik dikke grof fijn
takje takjes scheutje kneepje tl el cm gram kilo ongeveer sap rasp handje handjes bosje snufje
garnering garneren serveren topping zonder vel huid graat graten pit bot korst
rijpe koud koude warme zachte middelgrote flinke groot grote lange korte goede lekker
blad bos krop kropje bollen stronken plakken gesnipperd gehakt gesneden geschild gehalveerd
geperst gesmolten geplukt gekneusd geraspt geraspte gepeld uitgelekt losgeklopt schoongemaakt
afgekoeld fijngesneden fijngehakt grofgehakt grofgesneden gesnipperde
vers verse gedroogd gedroogde gemalen bio biologisch biologische klein kleine extra
""".split())


def kandidaten(rij: dict, mapping: dict[str, dict], met_product: dict[str, list[dict]]) -> list[tuple[str, str, str]]:
    """
    Voorstellen voor een regel zonder product: lijst van (mapping-sleutel, reden, sterkte).
    Sterkte: 'sterk' (zelfde lijstregel of keuze), 'redelijk' (deel van de naam), 'zwak' (lijkt op).
    """
    uit: list[tuple[str, str, str]] = []
    gezien: set[str] = set()

    def voeg(sleutel: str, reden: str, sterkte: str) -> None:
        if sleutel in mapping and sleutel not in gezien:
            gezien.add(sleutel)
            uit.append((sleutel, reden, sterkte))

    naam = rij['naam']
    # Een lege sleutel ("gehakt" valt als bereidingswoord weg): zoek dan op de platte naam.
    key = rij['ingredient_key'] or plat(naam)

    # 1. Een andere naam met dezelfde lijstregel heeft wél een product.
    for ander in met_product.get(rij['lijst_sleutel'], []):
        voeg(ander['ah_sleutel'], f'zelfde lijstregel als "{ander["naam"]}"', 'sterk')

    # 2. Een keuze: "pecorino of parmezaan", "tamari of sojasaus", "x (of y)".
    #    Een opsomming tussen haakjes zonder "of" ("groentemix (wortel, erwten, mais)")
    #    beschrijft wat erin zit en is maar een redelijk voorstel.
    delen: list[tuple[str, str]] = []
    haakjes = re.match(r'^([^(]*)\(([^)]*)\)', naam)
    if haakjes:
        sterkte = 'sterk' if re.search(r'\bof\b', haakjes.group(2)) else 'redelijk'
        delen += [(d, sterkte) for d in re.split(r',| of ', haakjes.group(2))]
        delen.append((haakjes.group(1), 'sterk'))
    if re.search(r'\bof\b|/', naam):
        delen += [(d, 'sterk') for d in re.split(r'\bof\b|/', plat(naam))]
    for deel, sterkte in delen:
        d = plat(re.sub(r'^(?:zoals|bijv\.?|bijvoorbeeld|bv\.?)\s+', '', deel.strip()))
        if not d or d == key:
            continue
        kaal = ' '.join(w for w in d.split() if w not in ONSCHULDIG)
        reden = (f'keuze in het recept: "{deel.strip()}"' if sterkte == 'sterk'
                 else f'tussen haakjes (opsomming): "{deel.strip()}"')
        for vorm in vormen(d) + (vormen(kaal) if kaal and kaal != d else []):
            voeg(vorm, reden, sterkte)
        for ander in met_product.get(d, []):
            voeg(ander['ah_sleutel'], f'{reden}; op de lijst "{ander["naam"]}"', sterkte)

    # 3. Een mapping-sleutel die als hele woorden in de naam zit; wat overblijft
    #    zegt waarom de app 'm niet nam (bv. "gerookte", "uit blik").
    woorden = key.split()
    beste: list[tuple[int, str, list[str]]] = []
    for sleutel in mapping:
        deel = sleutel.split()
        if not deel or len(deel) >= len(woorden):
            continue
        for i in range(len(woorden) - len(deel) + 1):
            if woorden[i:i + len(deel)] == deel:
                rest = woorden[:i] + woorden[i + len(deel):]
                beste.append((len(deel), sleutel, rest))
                break
    for _, sleutel, rest in sorted(beste, key=lambda b: -b[0])[:3]:
        blokkeert = [w for w in rest if w not in ONSCHULDIG]
        if blokkeert:
            voeg(sleutel, f'deel van de naam; "{" ".join(blokkeert)}" houdt de match tegen', 'redelijk')
        else:
            voeg(sleutel, 'deel van de naam', 'redelijk')

    # 4. Enkelvoud/meervoud en gelijkende sleutels.
    for vorm in vormen(key):
        voeg(vorm, 'enkel- of meervoud', 'redelijk')
    laatste = woorden[-1] if woorden else ''
    for sleutel in difflib.get_close_matches(key, list(mapping), n=4, cutoff=0.82):
        voeg(sleutel, 'lijkt op de sleutel', 'zwak')
    if laatste and len(laatste) >= 5:
        # De kortste eerst: "rundergehakt" vóór "jalapenoreepjes fijngehakt".
        for sleutel in sorted((s for s in mapping if s != key and s.split()[-1].endswith(laatste)), key=len):
            voeg(sleutel, f'eindigt ook op "{laatste}"', 'zwak')
    return uit[:6]


def vormen(key: str) -> list[str]:
    """Enkel- en meervoud van het laatste woord, zoals enkelvoudVormen in de app, plus simpele meervouden."""
    woorden = key.split()
    if not woorden:
        return []
    laatste = woorden[-1]
    voor = ' '.join(woorden[:-1]) + ' ' if len(woorden) > 1 else ''
    uit = [key]
    if len(laatste) >= 4 and laatste.endswith('s'):
        uit.append(voor + laatste[:-1])
    if len(laatste) >= 4 and laatste.endswith('en'):
        stam = laatste[:-2]
        uit.append(voor + stam)
        if len(stam) >= 2 and stam[-1] == stam[-2] and stam[-1] not in 'aeiou':
            uit.append(voor + stam[:-1])
        m = re.match(r'^(.*[^aeiou])([aeou])([^aeiou])$', stam)
        if m:
            uit.append(voor + m.group(1) + m.group(2) * 2 + m.group(3))
    uit += [voor + laatste + 's', voor + laatste + 'en']
    m = re.match(r'^(.*[^aeiou])([aeou])\2([^aeiou])$', laatste)
    if m:
        uit.append(voor + m.group(1) + m.group(2) + m.group(3) + 'en')
    return uit


def status(rij: dict) -> str:
    if rij['soort'] == 'altijd in huis':
        return 'Niet op de lijst (altijd in huis)'
    if rij['soort'] == 'toko':
        return 'Toko-link'
    if rij['soort'] in ('droog kruid', 'bijzonder kruid'):
        return 'Kruidenrek' + (' (AH-product)' if rij['ah_product'] else ' (zoeklink)')
    if rij['vega_variant'] and rij['vega_ah_product']:
        return f'Mandje als {rij["vega_variant"]} (vega-standaard)' + ('' if rij['ah_product'] else '; recept-keuze = zoeklink')
    if rij['ah_product']:
        return 'Mandje' if rij['ah_leverbaar'] is not False else 'Mandje, maar niet leverbaar'
    return 'Zoeklink'


def kopregel(ws, koppen: list[str], breedtes: list[int]) -> None:
    ws.append(koppen)
    for i, (kop, breedte) in enumerate(zip(koppen, breedtes), start=1):
        cel = ws.cell(row=1, column=i)
        cel.font = KOP
        cel.fill = KOPVULLING
        cel.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = breedte
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'B2'


def als_tabel(ws, naam: str, aantal_kolommen: int) -> None:
    if ws.max_row < 2:
        return
    ref = f'A1:{get_column_letter(aantal_kolommen)}{ws.max_row}'
    tabel = Table(displayName=naam, ref=ref)
    tabel.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
    ws.add_table(tabel)
    for rij in ws.iter_rows(min_row=2):
        for cel in rij:
            cel.font = Font(name=FONT)
            cel.alignment = Alignment(vertical='top', wrap_text=False)


def main() -> None:
    if not BRON.exists():
        sys.exit(f'{BRON.name} ontbreekt. Draai eerst: npm run ingredienten -- --offline')
    data = json.loads(BRON.read_text())
    ingredienten: list[dict] = data['ingredienten']
    mapping = {m['key']: m for m in data['mapping']}

    # Per lijstregel: welke namen krijgen een product?
    met_product: dict[str, list[dict]] = defaultdict(list)
    for r in ingredienten:
        if r['ah_product'] and r['soort'] == 'gewoon':
            met_product[r['lijst_sleutel']].append(r)

    wb = Workbook()

    # ---- Leeswijzer -------------------------------------------------------
    ws = wb.active
    ws.title = 'Leeswijzer'
    ws.column_dimensions['A'].width = 120
    gewoon = [r for r in ingredienten if r['soort'] == 'gewoon']
    zoek = [r for r in ingredienten if r['soort'] in ('gewoon', 'bijzonder kruid') and not r['ah_product']
            and not (r['vega_variant'] and r['vega_ah_product'])]
    regels_totaal = sum(r['aantal_regels'] for r in ingredienten)
    regels_zoek = sum(r['aantal_regels'] for r in zoek)
    regels = [
        ('Ingrediënten uit alle recepten en wat de app ermee doet', True),
        (f'Bron: {data["bron"]} · {data["aantal_recepten"]} recepten · gemaakt op {data["datum"]}', False),
        ('', False),
        ('Doel: namen samenvoegen zodat meer regels van de boodschappenlijst in het AH-mandje landen '
         'in plaats van als zoeklink. Een voorbeeld: een regel als "pecorino of parmezaan" matcht niets als '
         'geen van beide woorden een mapping-sleutel is; wordt het "parmezaanse kaas", dan gaat er een stuk '
         'parmezaan mee.', False),
        ('', False),
        ('Cijfers', True),
        (f'{len(ingredienten)} verschillende ingrediëntnamen in {regels_totaal} receptregels.', False),
        (f'{len(gewoon)} daarvan horen in het mandje (de rest: altijd in huis, kruidenrek, toko).', False),
        (f'{len(zoek)} namen krijgen nu een zoeklink, goed voor {regels_zoek} receptregels '
         f'({regels_zoek / regels_totaal:.0%} van alle regels).', False),
        (f'{len(mapping)} AH-mapping-regels, waarvan {sum(1 for m in mapping.values() if m["gebruik_regels"] == 0)} '
         'door geen enkele receptregel in deze bron geraakt worden (de database heeft ook je eigen '
         'imports; draai zonder --offline om die mee te nemen).', False),
        ('', False),
        ('Tabbladen', True),
        ('Zoeklinks — de actielijst. Elk ingrediënt dat nu een zoeklink krijgt, gesorteerd op hoe vaak het '
         'voorkomt. Kolom "Voorstel" is de mapping-sleutel waar het naartoe zou kunnen, met het AH-product '
         'dat daar aan hangt en de reden. Vul kolom "Besluit" in: de sleutel die het moet worden, '
         '"nieuw" als er een eigen mapping-regel voor moet komen, of "zoeklink" als het zo goed is.', False),
        ('Dubbele lijstregels — namen die hetzelfde AH-product krijgen maar op de lijst als aparte regels '
         'staan (andere lijstsleutel). Die horen één regel te worden via EEN_VORM of ANDERS_GESPELD in '
         'app/src/lib/synoniemen.ts.', False),
        ('Alle ingrediënten — elke naam, met ingredient_key, lijstsleutel, schap, status en het AH- en '
         'Jumbo-product. Filter op kolom Status.', False),
        ('AH-mapping — elke mapping-regel met het product, de leverbaarheid en hoe vaak hij geraakt wordt.', False),
        ('', False),
        ('Hoe de app zoekt (kort)', True),
        ('De naam uit het recept wordt een ingredient_key (kleine letters, zonder accenten, zonder "verse", '
         '"gesneden" enz.). De lijst voegt samen op de lijstsleutel (synoniemen.ts). Het product zoekt '
         'zoekProduct op: eerst de hele sleutel, dan enkel-/meervoud, dan een keuze ("x of y" → de eerste '
         'met een product), dan een deel van de naam, maar alleen als wat wegvalt het product niet '
         'verandert ("gerookte" en "uit blik" houden de match tegen).', False),
        ('', False),
        ('Sterkte van een voorstel', True),
        ('sterk = een andere spelling van hetzelfde ingrediënt heeft al een product, of het recept geeft zelf '
         'de keuze. redelijk = een deel van de naam is een bekend product; kijk of het weggevallen woord '
         'echt niets uitmaakt. zwak = lijkt er alleen op; vaak een ander product.', False),
        ('', False),
        ('Kleuren', True),
        ('Rood = zoeklink · Oranje = product bestaat maar is niet leverbaar · Lichtgeel = kolom om in te vullen.', False),
        ('', False),
        ('Let op: PR 116 (AH-mapping, 30 niet leverbare producten) stond nog open toen dit gemaakt werd; '
         'dit overzicht is op master gemaakt. Na die merge het script opnieuw draaien.', False),
    ]
    for tekst, vet in regels:
        ws.append([tekst])
        cel = ws.cell(row=ws.max_row, column=1)
        cel.font = Font(name=FONT, bold=vet, size=12 if vet else 10)
        cel.alignment = Alignment(wrap_text=True, vertical='top')

    # ---- Zoeklinks --------------------------------------------------------
    ws = wb.create_sheet('Zoeklinks')
    koppen = ['Ingrediënt (zoals in recept)', 'Recepten', 'Regels', 'ingredient_key', 'Lijstsleutel', 'Schap',
              'Keuze in recept', 'Voorstel (mapping-sleutel)', 'AH-product bij voorstel', 'Reden', 'Sterkte',
              'Andere mogelijkheden', 'Jumbo nu', 'Besluit (sleutel / nieuw / zoeklink)', 'Opmerking', 'Voorbeeldrecepten']
    kopregel(ws, koppen, [36, 9, 8, 28, 26, 18, 8, 28, 34, 40, 9, 44, 14, 30, 30, 60])
    for r in zoek:
        kand = kandidaten(r, mapping, met_product)
        # Een zwakke gelijkenis is vaak een ander product: die staat alleen bij de andere mogelijkheden.
        eerste = kand[0] if kand and kand[0][2] != 'zwak' else None
        rest = '; '.join(f'{k} ({m["naam"] or "?"}; {reden}, {sterkte})'
                         for k, reden, sterkte in (kand[1:] if eerste else kand) for m in [mapping[k]])
        ws.append([
            r['naam'], r['aantal_recepten'], r['aantal_regels'], r['ingredient_key'], r['lijst_sleutel'], r['schap'],
            'ja' if r['keuze'] else '',
            eerste[0] if eerste else '',
            (mapping[eerste[0]]['naam'] or '') if eerste else '',
            eerste[1] if eerste else ('alleen zwakke gelijkenissen, zie hiernaast' if kand else 'geen kandidaat gevonden'),
            eerste[2] if eerste else '',
            rest,
            r['jumbo_product'] or '',
            '', '',
            ' · '.join(r['recepten']),
        ])
    als_tabel(ws, 'Zoeklinks', len(koppen))
    for rij in ws.iter_rows(min_row=2, max_row=ws.max_row):
        rij[0].fill = ROOD
        rij[13].fill = INVUL
        rij[14].fill = INVUL
        sterkte = rij[10].value
        if sterkte == 'sterk':
            rij[7].fill = GROEN
        elif sterkte == 'redelijk':
            rij[7].fill = GEEL

    # ---- Dubbele lijstregels ---------------------------------------------
    ws = wb.create_sheet('Dubbele lijstregels')
    koppen = ['AH-productnummer', 'AH-product', 'Lijstsleutel', 'Namen in recepten', 'Recepten', 'Mapping-sleutel(s)']
    kopregel(ws, koppen, [16, 40, 30, 70, 10, 36])
    per_product: dict[str, dict[str, list[dict]]] = defaultdict(lambda: defaultdict(list))
    for r in gewoon:
        if r['ah_product']:
            per_product[str(r['ah_product'])][r['lijst_sleutel']].append(r)
    dubbel = [(p, s) for p, s in per_product.items() if len(s) > 1]
    dubbel.sort(key=lambda ps: -sum(r['aantal_recepten'] for rs in ps[1].values() for r in rs))
    for product, sleutels in dubbel:
        eerste = True
        for sleutel, rijen in sorted(sleutels.items(), key=lambda kv: -sum(r['aantal_recepten'] for r in kv[1])):
            ws.append([
                product if eerste else '',
                (rijen[0]['ah_naam'] or '') if eerste else '',
                sleutel,
                ' · '.join(sorted({r['naam'] for r in rijen}, key=str.lower)),
                sum(r['aantal_recepten'] for r in rijen),
                ', '.join(sorted({r['ah_sleutel'] for r in rijen})),
            ])
            eerste = False
        ws.append([''] * len(koppen))
    for rij in ws.iter_rows(min_row=2):
        for cel in rij:
            cel.font = Font(name=FONT, bold=bool(rij[0].value))
            cel.alignment = Alignment(vertical='top')

    # ---- Alle ingrediënten -------------------------------------------------
    ws = wb.create_sheet('Alle ingrediënten')
    koppen = ['Ingrediënt (zoals in recept)', 'Recepten', 'Regels', 'Voorbeeld hoeveelheid', 'Status', 'Soort', 'Schap',
              'ingredient_key', 'Lijstsleutel', 'AH mapping-sleutel', 'AH-nummer', 'AH-product', 'AH leverbaar', 'AH bio',
              'Vega-standaard', 'Jumbo-SKU', 'Jumbo-product', 'Voorbeeldrecepten']
    kopregel(ws, koppen, [36, 9, 8, 14, 30, 14, 18, 28, 26, 26, 11, 36, 10, 7, 14, 12, 36, 60])
    for r in ingredienten:
        ws.append([
            r['naam'], r['aantal_recepten'], r['aantal_regels'], r['voorbeeld_hoeveelheid'] or '',
            status(r), r['soort'], r['schap'], r['ingredient_key'], r['lijst_sleutel'],
            r['ah_sleutel'] or '', r['ah_product'] or '', r['ah_naam'] or '',
            '' if r['ah_leverbaar'] is None else ('ja' if r['ah_leverbaar'] else 'nee'),
            'ja' if r['ah_bio'] else '',
            r['vega_variant'] or '', r['jumbo_product'] or '', r['jumbo_naam'] or '',
            ' · '.join(r['recepten']),
        ])
    als_tabel(ws, 'AlleIngredienten', len(koppen))
    n = ws.max_row
    ws.conditional_formatting.add(f'A2:E{n}', FormulaRule(formula=['$E2="Zoeklink"'], fill=ROOD))
    ws.conditional_formatting.add(f'A2:E{n}', FormulaRule(formula=['$M2="nee"'], fill=ORANJE))

    # ---- AH-mapping ----------------------------------------------------------
    ws = wb.create_sheet('AH-mapping')
    koppen = ['Mapping-sleutel', 'AH-nummer', 'AH-product', 'Inhoud', 'Prijs', 'Leverbaar', 'Bio-nummer',
              'Jumbo-SKU', 'Jumbo-product', 'Geraakt (receptregels)', 'Door welke namen']
    kopregel(ws, koppen, [30, 11, 40, 10, 8, 10, 11, 12, 36, 12, 70])
    for m in sorted(mapping.values(), key=lambda m: (-m['gebruik_regels'], m['key'])):
        ws.append([
            m['key'], m['product'], m['naam'] or '', m['inhoud'] or '', m['prijs'],
            '' if m['leverbaar'] is None else ('ja' if m['leverbaar'] else 'nee'),
            m['bio'] or '', m['jumbo'] or '', m['jumbo_naam'] or '', m['gebruik_regels'],
            ' · '.join(sorted(set(m['gebruik_namen']), key=str.lower)),
        ])
    als_tabel(ws, 'AhMapping', len(koppen))
    n = ws.max_row
    ws.conditional_formatting.add(f'A2:C{n}', FormulaRule(formula=['$F2="nee"'], fill=ORANJE))
    ws.conditional_formatting.add(f'A2:C{n}', FormulaRule(formula=['$J2=0'], fill=GEEL))
    for rij in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        rij[0].number_format = '€ 0.00'

    wb.save(DOEL)
    sterk = sum(1 for r in zoek if (k := kandidaten(r, mapping, met_product)) and k[0][2] == 'sterk')
    redelijk = sum(1 for r in zoek if (k := kandidaten(r, mapping, met_product)) and k[0][2] == 'redelijk')
    print(f'{len(zoek)} zoeklinks ({regels_zoek} receptregels), {sterk} met een sterk en {redelijk} met een redelijk voorstel; '
          f'{len(dubbel)} producten met dubbele lijstregels.')
    print(f'Geschreven: {DOEL.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
